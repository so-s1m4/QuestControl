from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = "output/documents/Arbeitszeitaufzeichnung_Template.docx"
SEED = "apps/api/templates/Arbeitszeitaufzeichnung_Template.docx"

def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in kwargs[edge].items():
            element.set(qn("w:" + key), str(value))

def set_table_border(table, edge, size="6", color="000000"):
    """Set a table-level border, which survives merged header cells."""
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    element = borders.find(qn("w:" + edge))
    if element is None:
        element = OxmlElement("w:" + edge)
        borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), size)
    element.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, size=9, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(cell, top={"val":"single","sz":"6","color":"000000"}, bottom={"val":"single","sz":"6","color":"000000"}, left={"val":"single","sz":"6","color":"000000"}, right={"val":"single","sz":"6","color":"000000"})

def set_table_to_content_width(table, section):
    """Use the entire printable area, leaving only the page margins."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    width_twips = round((section.page_width - section.left_margin - section.right_margin) / 635)
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:w"), str(width_twips))
    table_width.set(qn("w:type"), "dxa")

def merge(table, start_row, start_col, end_row, end_col):
    return table.cell(start_row, start_col).merge(table.cell(end_row, end_col))

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.25)
section.left_margin = Cm(1.35)
section.right_margin = Cm(1.35)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(5)
title.paragraph_format.space_after = Pt(15)
run = title.add_run("Arbeitszeitaufzeichnung")
run.bold = True
run.font.name = "Arial"
run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
run.font.size = Pt(12)

meta = doc.add_table(rows=2, cols=2)
meta.autofit = False
set_table_to_content_width(meta, section)
for row in meta.rows:
    row.cells[0].width = Cm(7)
    row.cells[1].width = Cm(10.6)
set_cell_text(meta.cell(0, 0), "Name Arbeitnehmer/in:", False, 10, WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(meta.cell(0, 1), "{{employee_name}}", False, 10, WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(meta.cell(1, 0), "Monat: {{month}}", False, 10, WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(meta.cell(1, 1), "Jahr: {{year}}", False, 10, WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
table = doc.add_table(rows=33, cols=7)
table.autofit = False
set_table_to_content_width(table, section)
widths = [1.45, 2.35, 2.35, 2.35, 2.35, 2.85, 2.35]
for row in table.rows:
    for index, width in enumerate(widths):
        row.cells[index].width = Cm(width)

merge(table, 0, 0, 1, 0)
merge(table, 0, 1, 0, 2)
merge(table, 0, 3, 0, 4)
merge(table, 0, 5, 1, 5)
merge(table, 0, 6, 1, 6)
# python-docx drops parts of the top border when cells are merged. Restore it
# explicitly so the header remains a closed table in Word and in PDF export.
for col in range(7):
    set_cell_border(table.cell(0, col), top={"val":"single","sz":"6","color":"000000"})
set_table_border(table, "top")
set_cell_text(table.cell(0, 0), "Tag", True, 9)
set_cell_text(table.cell(0, 1), "Arbeitszeit", True, 9)
set_cell_text(table.cell(0, 3), "Pause", True, 9)
set_cell_text(table.cell(0, 5), "Tagesarbeitszeit\n(ohne Pause)", True, 8)
set_cell_text(table.cell(0, 6), "Notizen", True, 9)
for col, text in enumerate(["", "Beginn", "Ende", "Beginn", "Ende", "", ""]):
    if text:
        set_cell_text(table.cell(1, col), text, True, 9)

for day in range(1, 32):
    row = day + 1
    field = day - 1
    row_key = f"r{day:02d}"
    values = [
        f"{{{{{row_key}.d}}}}",
        f"{{{{{row_key}.s}}}}",
        f"{{{{{row_key}.e}}}}",
        f"{{{{{row_key}.ps}}}}",
        f"{{{{{row_key}.pe}}}}",
        f"{{{{{row_key}.h}}}}",
        f"{{{{{row_key}.n}}}}",
    ]
    for col, value in enumerate(values):
        set_cell_text(table.cell(row, col), value, False, 8 if col == 6 else 9, WD_ALIGN_PARAGRAPH.LEFT if col == 6 else WD_ALIGN_PARAGRAPH.RIGHT)
    table.rows[row].height = Cm(0.54)

doc.add_paragraph().paragraph_format.space_after = Pt(3)
signatures = doc.add_table(rows=2, cols=2)
signatures.autofit = False
set_table_to_content_width(signatures, section)
for row in signatures.rows:
    row.cells[0].width = Cm(8.6)
    row.cells[1].width = Cm(8.6)
set_cell_text(signatures.cell(0, 0), "________________________\n{{employee_signature_date}}", False, 9, WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(signatures.cell(0, 1), "________________________\n{{employer_signature_date}}", False, 9, WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(signatures.cell(1, 0), "{{employee_signature}}\nUnterschrift Arbeitnehmer/in", False, 9, WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(signatures.cell(1, 1), "{{employer_signature}}\nUnterschrift Arbeitgeber/in", False, 9, WD_ALIGN_PARAGRAPH.LEFT)

doc.core_properties.title = "Arbeitszeitaufzeichnung"
doc.core_properties.subject = "Vorlage für monatliche Arbeitszeitaufzeichnungen"
doc.save(OUT)
Path(SEED).parent.mkdir(parents=True, exist_ok=True)
doc.save(SEED)
