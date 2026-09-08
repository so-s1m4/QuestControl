import base64
import json
import os
import re
import resource
import subprocess
import sys
import tempfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches

MAX_BODY = 20_000_000
MAX_SCRIPT = 300_000
MAX_OUTPUT = 2_000_000
PLACEHOLDER = re.compile(r"{{\s*([A-Za-z_][A-Za-z0-9_.]*)\s*}}")
SOFFICE_BINARY = os.environ.get("SOFFICE_BINARY", "soffice")

HARNESS = '''import json, runpy, sys
context = json.loads(sys.stdin.read())
module = runpy.run_path(sys.argv[1], run_name="user_document_script")
run = module.get("run")
if not callable(run):
    raise RuntimeError("Python script must define run(context)")
result = run(context)
if not isinstance(result, dict):
    raise RuntimeError("run(context) must return a JSON object")
print(json.dumps(result, ensure_ascii=False, default=str))
'''

def child_limits():
    limits = [
        (resource.RLIMIT_CPU, 5, 6),
        (resource.RLIMIT_AS, 256 * 1024 * 1024, 256 * 1024 * 1024),
        (resource.RLIMIT_FSIZE, 12 * 1024 * 1024, 12 * 1024 * 1024),
        (resource.RLIMIT_NOFILE, 32, 32),
        (getattr(resource, "RLIMIT_NPROC", None), 12, 12),
    ]
    for limit, soft, hard in limits:
        if limit is None:
            continue
        try:
            resource.setrlimit(limit, (soft, hard))
        except (ValueError, OSError):
            pass

def run_script(source_code, context, workdir):
    if len(source_code.encode("utf-8")) > MAX_SCRIPT:
        raise ValueError("Script is too large")
    script = Path(workdir) / "script.py"
    harness = Path(workdir) / "harness.py"
    script.write_text(source_code, encoding="utf-8")
    harness.write_text(HARNESS, encoding="utf-8")
    result = subprocess.run(
        [sys.executable, "-I", str(harness), str(script)],
        input=json.dumps(context, ensure_ascii=False), text=True, cwd=workdir,
        env={"PATH": "/usr/local/bin:/usr/bin:/bin", "PYTHONDONTWRITEBYTECODE": "1"},
        capture_output=True, timeout=7, preexec_fn=child_limits,
    )
    if result.returncode != 0:
        detail = (result.stderr or result.stdout or "Script failed").strip()[-600:]
        raise ValueError(detail)
    if len(result.stdout.encode("utf-8")) > MAX_OUTPUT:
        raise ValueError("Script result is too large")
    try:
        return json.loads(result.stdout)
    except json.JSONDecodeError as error:
        raise ValueError("Script must print a JSON object only") from error

def resolve(data, path):
    value = data
    for part in path.split("."):
        if not isinstance(value, dict) or part not in value:
            return ""
        value = value[part]
    return value

def all_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from all_paragraphs(cell)

def add_centered_signature_overlay(paragraph, image_path):
    """Place a signature over its caption instead of consuming a text line."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), height=Inches(0.42))

    # Word represents regular pictures as ``wp:inline``. Converting that one
    # drawing to a floating anchor lets the transparent signature sit above the
    # caption ("Unterschrift …") while the caption stays selectable text.
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "251659264", "behindDoc": "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    position_h.append(align_h)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "line")
    offset_v = OxmlElement("wp:posOffset")
    # Shift down from the marker line onto the caption below it.
    offset_v.text = "95000"
    position_v.append(offset_v)
    wrap_none = OxmlElement("wp:wrapNone")
    inline_xml = inline._inline
    anchor.extend([simple_pos, position_h, position_v, inline_xml.extent, wrap_none, inline_xml.docPr, inline_xml.graphic])
    drawing = inline_xml.getparent()
    drawing.replace(inline_xml, anchor)

def replace_in_paragraph(paragraph, values, temp_dir):
    text = paragraph.text
    if not PLACEHOLDER.search(text):
        return
    image_matches = []
    for match in PLACEHOLDER.finditer(text):
        value = resolve(values, match.group(1))
        if isinstance(value, dict) and value.get("_type") == "image":
            image_matches.append((match, value))
    if image_matches:
        # Signature tags are deliberately kept alone in their paragraph in the supplied template.
        match, image = image_matches[0]
        before, after = text[:match.start()], text[match.end():]
        content_type = image.get("content_type", "image/png")
        extension = ".jpg" if content_type == "image/jpeg" else ".png"
        image_path = Path(temp_dir) / ("signature" + extension)
        image_path.write_bytes(base64.b64decode(image["data_base64"], validate=True))
        for run in paragraph.runs:
            run.text = ""
        paragraph.add_run(before)
        add_centered_signature_overlay(paragraph, image_path)
        paragraph.add_run(after)
        return
    rendered = PLACEHOLDER.sub(lambda match: str(resolve(values, match.group(1)) or ""), text)
    if paragraph.runs:
        paragraph.runs[0].text = rendered
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(rendered)

def render_docx(template_data, values, temp_dir):
    template_path = Path(temp_dir) / "template.docx"
    output_path = Path(temp_dir) / "result.docx"
    template_path.write_bytes(template_data)
    doc = Document(str(template_path))
    for paragraph in all_paragraphs(doc):
        replace_in_paragraph(paragraph, values, temp_dir)
    doc.save(str(output_path))
    return output_path.read_bytes()

def convert_docx_to_pdf(document_data, temp_dir):
    source_path = Path(temp_dir) / "generated.docx"
    output_dir = Path(temp_dir) / "pdf"
    profile_dir = Path(temp_dir) / "office-profile"
    output_dir.mkdir()
    profile_dir.mkdir()
    source_path.write_bytes(document_data)
    result = subprocess.run(
        [SOFFICE_BINARY, f"-env:UserInstallation=file://{profile_dir}", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(source_path)],
        capture_output=True, text=True, timeout=12, env={**os.environ, "HOME": temp_dir, "TMPDIR": temp_dir},
    )
    output_path = output_dir / "generated.pdf"
    if result.returncode != 0 or not output_path.exists():
        raise ValueError("Unable to convert DOCX to PDF")
    document = output_path.read_bytes()
    if not document.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF output")
    return document

class Handler(BaseHTTPRequestHandler):
    def send_json(self, status, data):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        if self.path == "/health":
            return self.send_json(200, {"ok": True})
        self.send_json(404, {"error": "NOT_FOUND"})

    def do_POST(self):
        if self.path != "/generate":
            return self.send_json(404, {"error": "NOT_FOUND"})
        try:
            size = int(self.headers.get("Content-Length", "0"))
            if size <= 0 or size > MAX_BODY:
                raise ValueError("Invalid request size")
            payload = json.loads(self.rfile.read(size))
            source = payload["sourceCode"]
            context = payload["context"]
            defaults = payload.get("defaults", {})
            output_format = payload.get("outputFormat", "docx")
            template_data = base64.b64decode(payload["templateDataBase64"], validate=True)
            if not isinstance(source, str) or not isinstance(context, dict) or not isinstance(defaults, dict) or output_format not in {"docx", "pdf"}:
                raise ValueError("Invalid request")
            if len(template_data) > 8_000_000 or not template_data.startswith(b"PK"):
                raise ValueError("Invalid DOCX template")
            with tempfile.TemporaryDirectory(prefix="document-") as workdir:
                result = run_script(source, context, workdir)
                values = {**defaults, **result}
                document = render_docx(template_data, values, workdir)
                if output_format == "pdf":
                    document = convert_docx_to_pdf(document, workdir)
            self.send_json(200, {"documentBase64": base64.b64encode(document).decode("ascii"), "outputFormat": output_format})
        except subprocess.TimeoutExpired:
            self.send_json(422, {"error": "SCRIPT_TIMEOUT", "message": "Script exceeded the 7 second limit"})
        except (ValueError, KeyError, UnicodeDecodeError, OSError) as error:
            self.send_json(422, {"error": "DOCUMENT_GENERATION_FAILED", "message": str(error)[:700]})
        except Exception:
            self.send_json(500, {"error": "DOCUMENT_GENERATION_FAILED", "message": "Unable to generate document"})

    def log_message(self, _format, *_args):
        pass

if __name__ == "__main__":
    ThreadingHTTPServer(("0.0.0.0", 8080), Handler).serve_forever()
